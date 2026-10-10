#!/usr/bin/env python3
"""Local server for Patent Workbench modules 1, 2 and 3."""
from __future__ import annotations

import io
import json
import os
import re
import sys
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import quote, unquote, urlsplit
from urllib.request import Request, urlopen

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent
MAX_BODY = 5 * 1024 * 1024
PROMPT_FILE = ROOT / "docs" / "특허명세서_마스터_프롬프트.md"


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(ROOT), **kwargs)

    def log_message(self, fmt, *args):
        # Do not write request body content, which can contain confidential invention details.
        sys.stderr.write("%s - %s\n" % (self.address_string(), fmt % args))

    def do_OPTIONS(self):
        self.send_error(405)

    def do_GET(self):
        if self.path == "/api/status":
            self._json(200, {"configured": bool(os.environ.get("OPENAI_API_KEY")), "model": os.environ.get("OPENAI_MODEL", "gpt-5-mini"), "module_version": "0.2.0", "packet_schema_version": "1.0"})
            return
        # The runtime prompt is consumed by the server and is never an app route.
        route = unquote(urlsplit(self.path).path)
        if route in ("/docs/특허명세서_마스터_프롬프트.md", "/docs/특허명세서_마스터_프롬프트.docx"):
            self.send_error(404)
            return
        if self.path == "/":
            self.path = "/index.html"
        super().do_GET()

    def do_POST(self):
        if self.path not in ("/api/draft", "/api/revise", "/api/docx"):
            self.send_error(404)
            return
        try:
            payload = self._read_json()
            if self.path == "/api/draft":
                self._draft(payload)
            elif self.path == "/api/revise":
                self._revise(payload)
            else:
                self._docx(payload)
        except ValueError as exc:
            self._json(400, {"error": str(exc)})
        except Exception as exc:
            self._json(500, {"error": str(exc) or "요청을 처리하지 못했습니다."})

    def _read_json(self):
        try:
            length = int(self.headers.get("Content-Length", "0"))
        except ValueError:
            raise ValueError("요청 크기를 확인할 수 없습니다.")
        if length <= 0 or length > MAX_BODY:
            raise ValueError("요청 본문이 비어 있거나 5MB를 초과했습니다.")
        try:
            return json.loads(self.rfile.read(length).decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError):
            raise ValueError("JSON 형식의 요청이 필요합니다.")

    def _json(self, status, value):
        data = json.dumps(value, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(data)

    def _draft(self, payload):
        packet = payload.get("packet") or {}
        if packet.get("packet_schema_version") != "1.0":
            raise ValueError("패킷 계약 버전 1.0 파일을 가져와 주세요.")
        claims = packet.get("claims", {}).get("text", "").strip()
        if not claims:
            raise ValueError("전달 패킷에 청구항 원문이 없습니다.")
        if not PROMPT_FILE.is_file():
            raise ValueError("서버의 마스터 프롬프트 문서를 찾을 수 없습니다.")
        key = os.environ.get("OPENAI_API_KEY", "").strip()
        if not key:
            raise ValueError("AI 서버에 OPENAI_API_KEY가 설정되지 않았습니다. README의 설정 방법을 확인해 주세요.")
        prompt = PROMPT_FILE.read_text(encoding="utf-8")
        user_payload = {
            "project": packet.get("project", {}),
            "claims": packet.get("claims", {}),
            "materials": packet.get("materials", []),
            "claim_components": packet.get("claim_components", []),
            "author_confirmations": packet.get("author_confirmations", []),
            "drafting_constraints": packet.get("drafting_constraints", {}),
            "user_additional_instruction": payload.get("extra_instruction", ""),
            "output_requirements": "한국어 특허명세서의 발명의 상세한 설명을 작성합니다. 청구항 원문은 변경하지 않습니다. 본문 뒤에 [A. 발명자 확인 필요사항], [B. 추가 실시예/Disclosure 확보 권고], [C. 청구항 관련 검토사항]을 구분해 둡니다. 기술적으로 의미 있는 소제목과 문단을 사용합니다."
        }
        model = os.environ.get("OPENAI_MODEL", "").strip() or str(payload.get("model", "gpt-5-mini")).strip()
        body = json.dumps({
            "model": model,
            "instructions": prompt,
            "input": [{"role": "user", "content": [{"type": "input_text", "text": json.dumps(user_payload, ensure_ascii=False)}]}],
            "max_output_tokens": int(os.environ.get("OPENAI_MAX_OUTPUT_TOKENS", "16000"))
        }, ensure_ascii=False).encode("utf-8")
        request = Request("https://api.openai.com/v1/responses", data=body, headers={"Authorization": "Bearer " + key, "Content-Type": "application/json"}, method="POST")
        try:
            with urlopen(request, timeout=300) as response:
                result = json.loads(response.read().decode("utf-8"))
        except HTTPError as exc:
            detail = exc.read().decode("utf-8", errors="replace")[:1200]
            raise ValueError(f"AI API 요청 실패 ({exc.code}): {detail}")
        except URLError as exc:
            raise ValueError(f"AI 서비스에 연결하지 못했습니다: {exc.reason}")
        text = result.get("output_text", "").strip()
        if not text:
            parts = []
            for item in result.get("output", []):
                for content in item.get("content", []):
                    if content.get("type") in ("output_text", "text"):
                        parts.append(content.get("text", ""))
            text = "\n".join(parts).strip()
        if not text:
            raise ValueError("AI 응답에서 본문 텍스트를 확인하지 못했습니다.")
        self._json(200, {"text": text, "model": model, "response_id": result.get("id", ""), "module_version": "0.2.0", "packet_schema_version": "1.0"})

    @staticmethod
    def _response_text(result):
        text = result.get("output_text", "").strip()
        if text:
            return text
        parts = []
        for item in result.get("output", []):
            for content in item.get("content", []):
                if content.get("type") in ("output_text", "text"):
                    parts.append(content.get("text", ""))
        return "\n".join(parts).strip()

    def _revise(self, payload):
        document = payload.get("document") or {}
        if document.get("document_schema_version") != "1.0":
            raise ValueError("문서 스키마 버전 1.0 초안이 필요합니다.")
        paragraphs = document.get("paragraphs") or []
        if not isinstance(paragraphs, list) or not paragraphs:
            raise ValueError("문서에 수정 가능한 문단이 없습니다.")
        instruction = str(payload.get("instruction", "")).strip()
        if not instruction:
            raise ValueError("수정 지시를 입력해 주세요.")
        raw_numbers = payload.get("paragraph_numbers") or []
        if not isinstance(raw_numbers, list) or not raw_numbers:
            raise ValueError("수정할 문단 번호를 하나 이상 지정해 주세요.")
        if any(not isinstance(n, int) or isinstance(n, bool) for n in raw_numbers):
            raise ValueError("문단 번호는 양의 정수여야 합니다.")
        numbers = sorted(set(raw_numbers))
        if any(n < 1 or n > len(paragraphs) for n in numbers):
            raise ValueError("문단 번호가 현재 문서 범위를 벗어났습니다.")
        key = os.environ.get("OPENAI_API_KEY", "").strip()
        if not key:
            raise ValueError("AI 서버에 OPENAI_API_KEY가 설정되지 않았습니다.")
        if not PROMPT_FILE.is_file():
            raise ValueError("서버의 마스터 프롬프트 문서를 찾을 수 없습니다.")
        model = os.environ.get("OPENAI_MODEL", "").strip() or str(payload.get("model", "gpt-5-mini")).strip()
        schema = {
            "type": "object", "properties": {"revisions": {"type": "array", "items": {
                "type": "object", "properties": {"paragraph_number": {"type": "integer"}, "content": {"type": "string"}},
                "required": ["paragraph_number", "content"], "additionalProperties": False
            }}}, "required": ["revisions"], "additionalProperties": False
        }
        revision_rules = (
            "문단별 수정 모듈의 제안 작성자다. 아래 마스터 프롬프트와 자료 근거를 따른다. "
            "지정 문단만 수정하고 나머지는 응답에 포함하지 않는다. 근거가 없는 새 기술 사실·수치·효과를 만들지 않는다. "
            "사용자가 요구한 문단 번호마다 수정된 문단 전체를 반환한다. 확인할 수 없는 사실은 [발명자 확인 필요]로 남긴다. "
            "청구항은 변경하지 않는다. 출력은 지정된 JSON schema에 맞춘다."
        )
        content = {
            "project": document.get("project", {}), "claims": document.get("claims", ""),
            "all_paragraphs": paragraphs, "paragraph_numbers_to_revise": numbers,
            "user_revision_instruction": instruction
        }
        request_body = json.dumps({
            "model": model,
            "instructions": PROMPT_FILE.read_text(encoding="utf-8") + "\n\n" + revision_rules,
            "input": [{"role": "user", "content": [{"type": "input_text", "text": json.dumps(content, ensure_ascii=False)}]}],
            "text": {"format": {"type": "json_schema", "name": "paragraph_revisions", "strict": True, "schema": schema}},
            "max_output_tokens": 8000
        }, ensure_ascii=False).encode("utf-8")
        request = Request("https://api.openai.com/v1/responses", data=request_body, headers={"Authorization": "Bearer " + key, "Content-Type": "application/json"}, method="POST")
        try:
            with urlopen(request, timeout=180) as response:
                result = json.loads(response.read().decode("utf-8"))
        except HTTPError as exc:
            detail = exc.read().decode("utf-8", errors="replace")[:1200]
            raise ValueError(f"AI 수정 요청 실패 ({exc.code}): {detail}")
        except URLError as exc:
            raise ValueError(f"AI 서비스에 연결하지 못했습니다: {exc.reason}")
        response_text = self._response_text(result)
        try:
            parsed = json.loads(response_text)
            revisions = parsed.get("revisions", [])
        except (json.JSONDecodeError, AttributeError):
            raise ValueError("AI 응답의 문단 수정 형식을 확인할 수 없습니다.")
        allowed = set(numbers)
        seen = set()
        for item in revisions:
            number = item.get("paragraph_number")
            if not isinstance(number, int) or isinstance(number, bool) or number not in allowed or number in seen or not isinstance(item.get("content"), str):
                raise ValueError("AI 응답의 문단 번호 또는 수정 내용이 유효하지 않습니다.")
            seen.add(number)
        if seen != allowed:
            raise ValueError("일부 지정 문단의 수정 제안을 받지 못했습니다. 다시 요청해 주세요.")
        revisions.sort(key=lambda item: item["paragraph_number"])
        self._json(200, {"revisions": revisions, "model": model, "response_id": result.get("id", ""), "module_version": "0.3.0", "document_schema_version": "1.0"})

    def _docx(self, payload):
        text = str(payload.get("text", "")).strip()
        if not text:
            raise ValueError("DOCX로 내보낼 본문이 없습니다.")
        document = Document()
        section = document.sections[0]
        section.top_margin = Cm(2.3)
        section.bottom_margin = Cm(2.3)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        normal = document.styles["Normal"]
        normal.font.name = "맑은 고딕"
        normal.font.size = Pt(10)
        normal.paragraph_format.line_spacing = 1.65
        title = document.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("발명의 상세한 설명")
        run.bold = True
        run.font.name = "맑은 고딕"
        run.font.size = Pt(17)
        document.add_paragraph("")
        for raw_line in text.splitlines():
            line = raw_line.strip()
            if not line:
                continue
            clean = re.sub(r"^#{1,6}\s*", "", line)
            clean = re.sub(r"^\*\*(.+?)\*\*$", r"\1", clean)
            if re.match(r"^(\[A\.|\[B\.|\[C\.|[0-9]+[.)、]\s|발명의 명칭|기술분야|배경기술|해결하려는 과제|과제의 해결 수단|발명의 효과|도면의 간단한 설명|부호의 설명|실시예|구체적인 내용)", clean):
                para = document.add_paragraph()
                para.paragraph_format.space_before = Pt(8)
                para.paragraph_format.space_after = Pt(4)
                run = para.add_run(clean)
                run.bold = True
                run.font.size = Pt(11)
            else:
                document.add_paragraph(clean)
        output = io.BytesIO()
        document.save(output)
        data = output.getvalue()
        filename = re.sub(r"[^\w가-힣.-]+", "_", str(payload.get("title", "특허명세서"))).strip("._")[:80] or "특허명세서"
        self.send_response(200)
        self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        version = str(payload.get("version", "")).strip()
        suffix = f"-v{version}" if version.isdigit() else ""
        self.send_header("Content-Disposition", f"attachment; filename*=UTF-8''{quote(filename + '_명세서' + suffix + '.docx')}")
        self.send_header("Content-Length", str(len(data)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(data)


def main():
    host = os.environ.get("PATENT_HOST", "127.0.0.1")
    port = int(os.environ.get("PATENT_PORT", "8000"))
    server = ThreadingHTTPServer((host, port), Handler)
    print(f"Patent Workbench: http://{host}:{port}/")
    print("AI generation is enabled when OPENAI_API_KEY is set in the server environment.")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nServer stopped.")
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
